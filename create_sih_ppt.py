from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_ppt():
    prs = Presentation()
    
    # Common helper for title slide
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    
    # --- Slide 1: Title ---
    slide1 = prs.slides.add_slide(title_slide_layout)
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    
    title.text = "SMART INDIA HACKATHON 2026"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    title.text_frame.paragraphs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = (
        "AI-Based Cognitive Gaming and Memory Assistance Platform for Elderly\n"
        "Dementia Patients in North Eastern Region (NER)\n\n"
        "• Problem Statement ID – 26003\n"
        "• Problem Statement Title- AI-Based Cognitive Gaming\n"
        "• Theme- MedTech / BioTech / HealthTech\n"
        "• PS Category- Software/Hardware\n"
        "• Team ID- [Your Team ID]\n"
        "• Team Name- Vectors"
    )
    for p in subtitle.text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        p.font.size = Pt(18)
        
    # --- Slide 2: Idea Description ---
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    title2 = slide2.shapes.title
    title2.text = "Culturally and Regionally Personalized Cognitive Support Platform"
    body2 = slide2.placeholders[1]
    tf2 = body2.text_frame
    tf2.text = "CARE - Caring And Remembering Everyday"
    
    p = tf2.add_paragraph()
    p.text = "An interactive, offline-first Progressive Web App (PWA) tailored for dementia patients in the NER."
    p.level = 1
    
    p = tf2.add_paragraph()
    p.text = "Delivers culturally familiar cognitive activities (e.g., Festival Faces) to trigger positive reminiscence."
    p.level = 1
    
    p = tf2.add_paragraph()
    p.text = "Utilizes localized text-to-speech (TTS) engines for Hindi, Assamese, and English without internet."
    p.level = 1

    # --- Slide 3: Technical Approach (Architecture) ---
    slide3 = prs.slides.add_slide(bullet_slide_layout)
    title3 = slide3.shapes.title
    title3.text = "TECHNICAL APPROACH (Architecture)"
    body3 = slide3.placeholders[1]
    tf3 = body3.text_frame
    
    tf3.text = "Technologies & Frameworks:"
    
    p = tf3.add_paragraph()
    p.text = "[Frontend Interface]: Next.js 16 (React), Tailwind CSS (Elder-friendly UI)"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "[Offline Engine (PWA)]: Service Workers, IndexedDB for local caching and offline telemetry."
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "[Native Interactions]: Web Speech API & Web Audio API (Procedural soundscapes)."
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "[Deployment]: Vercel (Hosting), GitHub (Version Control)."
    p.level = 1

    # --- Slide 4: Technical Approach (Methodology) ---
    slide4 = prs.slides.add_slide(bullet_slide_layout)
    title4 = slide4.shapes.title
    title4.text = "TECHNICAL APPROACH (Methodology)"
    body4 = slide4.placeholders[1]
    tf4 = body4.text_frame
    
    tf4.text = "Implementation Workflow:"
    
    p = tf4.add_paragraph()
    p.text = "1. Caregiver Setup: Selects patient's preferred language and cultural background."
    p.level = 1
    
    p = tf4.add_paragraph()
    p.text = "2. Offline Caching: PWA silently downloads assets for network independence."
    p.level = 1
    
    p = tf4.add_paragraph()
    p.text = "3. Cognitive Engagement: Patient plays Festival Faces (emotion ID) & Soundscape Painter."
    p.level = 1
    
    p = tf4.add_paragraph()
    p.text = "4. Telemetry Logging: Silently logs interaction times and accuracy locally."
    p.level = 1

    # --- Slide 5: Feasibility and Viability ---
    slide5 = prs.slides.add_slide(bullet_slide_layout)
    title5 = slide5.shapes.title
    title5.text = "FEASIBILITY AND VIABILITY"
    body5 = slide5.placeholders[1]
    tf5 = body5.text_frame
    
    tf5.text = "TECHNICAL FEASIBILITY:"
    p = tf5.add_paragraph()
    p.text = "Built on modern web standards; runs seamlessly as a PWA on low-end smartphones."
    p.level = 1
    
    p = tf5.add_paragraph()
    p.text = "CHALLENGES & RISKS:"
    p.level = 0
    p2 = tf5.add_paragraph()
    p2.text = "Device fragmentation (missing native TTS voices); Elderly UX friction."
    p2.level = 1
    
    p3 = tf5.add_paragraph()
    p3.text = "MITIGATION STRATEGIES:"
    p3.level = 0
    p4 = tf5.add_paragraph()
    p4.text = "Graceful degradation to Hindi/English text; touch-manipulation CSS to block zooming."
    p4.level = 1

    # --- Slide 6: Impact and Benefits ---
    slide6 = prs.slides.add_slide(bullet_slide_layout)
    title6 = slide6.shapes.title
    title6.text = "IMPACT AND BENEFITS"
    body6 = slide6.placeholders[1]
    tf6 = body6.text_frame
    
    tf6.text = "Impact on Elderly Dementia Patients in NER:"
    
    p = tf6.add_paragraph()
    p.text = "Cultural Familiarity: Bridges the gap with specific NER contexts (Bihu, Hornbill)."
    p.level = 1
    
    p = tf6.add_paragraph()
    p.text = "Social Inclusion: Provides digital therapy to those who only speak regional languages."
    p.level = 1
    
    p = tf6.add_paragraph()
    p.text = "Caregiver Relief: Safe, independent activity reduces constant monitoring strain."
    p.level = 1
    
    p = tf6.add_paragraph()
    p.text = "Economic Scalability: Zero distribution cost via PWA; deploys across iOS/Android instantly."
    p.level = 1

    # --- Slide 7: Research and References ---
    slide7 = prs.slides.add_slide(bullet_slide_layout)
    title7 = slide7.shapes.title
    title7.text = "RESEARCH AND REFERENCES"
    body7 = slide7.placeholders[1]
    tf7 = body7.text_frame
    
    tf7.text = "Project Links:"
    
    p = tf7.add_paragraph()
    p.text = "GitHub Repository: https://github.com/flyingmeai01-ship-it/SIH_Vector"
    p.level = 1
    
    p = tf7.add_paragraph()
    p.text = "Live Platform: https://sih-vector-psi.vercel.app"
    p.level = 1
    
    p = tf7.add_paragraph()
    p.text = "References:"
    p.level = 0
    
    p = tf7.add_paragraph()
    p.text = "W3C Web Content Accessibility Guidelines (WCAG) 2.1 for Elderly Users."
    p.level = 1
    
    prs.save("SIH_2026_CARE_Presentation.pptx")

if __name__ == "__main__":
    create_ppt()
